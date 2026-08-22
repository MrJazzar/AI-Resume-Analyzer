import json
from io import BytesIO
from docx import Document
import httpx
import pytest
from app.config import settings


def make_docx() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("John Doe")
    document.add_paragraph("Skills: Python, FastAPI, SQL, Git")
    document.save(output)
    return output.getvalue()


def test_full_end_to_end_flow(client, monkeypatch):
    # 1. Register
    reg_resp = client.post(
        "/api/auth/register",
        json={"name": "Alice", "email": "alice@example.com", "password": "SecurePass123"},
    )
    assert reg_resp.status_code == 201
    user_data = reg_resp.json()
    assert user_data["email"] == "alice@example.com"
    assert "session_id" in reg_resp.cookies

    # 2. Login
    # Since registration already logs in, let's clear cookies and test login explicitly
    client.cookies.clear()
    login_resp = client.post(
        "/api/auth/login",
        json={"email": "alice@example.com", "password": "SecurePass123"},
    )
    assert login_resp.status_code == 200
    assert "session_id" in login_resp.cookies

    # 3. Upload Resume
    upload_resp = client.post(
        "/api/resumes/upload",
        files={
            "file": (
                "resume.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload_resp.status_code == 201
    resume_data = upload_resp.json()
    resume_id = resume_data["id"]
    assert resume_data["filename"] == "resume.docx"
    assert resume_data["raw_text"] != ""

    # 4. Resume Analysis (Mocking the AI/Gemini endpoint)
    # Monkeypatch the analyzer service inside app.routes.resumes to bypass external API
    from app.routes import resumes
    monkeypatch.setattr(
        resumes,
        "analyze_resume",
        lambda text: {
            "summary": "Fullstack python programmer",
            "technical_skills": ["Python", "FastAPI", "SQL", "Git"],
            "soft_skills": ["Communication"],
            "education": [],
            "experience": [],
            "projects": [],
        },
    )

    analyze_resp = client.post(f"/api/resumes/{resume_id}/analyze")
    assert analyze_resp.status_code == 200
    analysis_data = analyze_resp.json()
    assert "Python" in analysis_data["technical_skills"]
    assert "FastAPI" in analysis_data["technical_skills"]

    # 5. Create Jobs (Authenticated)
    # Job 1: Backend Developer (Python, FastAPI, SQL) -> 100% Match (3 out of 3)
    j1_resp = client.post(
        "/api/jobs/",
        json={
            "title": "Backend Developer",
            "company": "TechCorp",
            "required_skills": ["Python", "FastAPI", "SQL"],
        },
    )
    assert j1_resp.status_code == 201
    j1_data = j1_resp.json()

    # Job 2: AI Engineer (Python, Machine Learning, TensorFlow) -> 33% Match (1 out of 3)
    j2_resp = client.post(
        "/api/jobs/",
        json={
            "title": "AI Engineer",
            "company": "AILab",
            "required_skills": ["Python", "Machine Learning", "TensorFlow"],
        },
    )
    assert j2_resp.status_code == 201
    j2_data = j2_resp.json()

    # Job 3: Frontend Developer (React, JavaScript, CSS) -> 0% Match (0 out of 3)
    j3_resp = client.post(
        "/api/jobs/",
        json={
            "title": "Frontend Developer",
            "company": "DesignStudio",
            "required_skills": ["React", "JavaScript", "CSS"],
        },
    )
    assert j3_resp.status_code == 201
    j3_data = j3_resp.json()

    # 6. Request Recommendations
    recs_resp = client.get(f"/api/recommendations/{resume_id}")
    assert recs_resp.status_code == 200
    recommendations = recs_resp.json()
    assert len(recommendations) == 3

    # Check Descending Order by Score: Job 1 (100%), Job 2 (33%), Job 3 (0%)
    assert recommendations[0]["job_id"] == j1_data["id"]
    assert recommendations[0]["score"] == 100
    assert recommendations[0]["matched_skills"] == ["FastAPI", "Python", "SQL"]
    assert recommendations[0]["missing_skills"] == []

    assert recommendations[1]["job_id"] == j2_data["id"]
    assert recommendations[1]["score"] == 33
    assert recommendations[1]["matched_skills"] == ["Python"]
    assert recommendations[1]["missing_skills"] == ["Machine Learning", "TensorFlow"]

    assert recommendations[2]["job_id"] == j3_data["id"]
    assert recommendations[2]["score"] == 0
    assert recommendations[2]["matched_skills"] == []
    assert recommendations[2]["missing_skills"] == ["CSS", "JavaScript", "React"]

    # 7. Request Recommendation Detail (Mocking the explanation)
    # Monkeypatch the API key and httpx.post inside app.services.job_matcher
    monkeypatch.setattr(settings, "AI_API_KEY", "mock-key")
    monkeypatch.setattr(
        httpx,
        "post",
        lambda *args, **kwargs: type(
            "MockResponse",
            (),
            {
                "status_code": 200,
                "raise_for_status": lambda self: None,
                "json": lambda self: {
                    "candidates": [
                        {
                            "content": {
                                "parts": [
                                    {
                                        "text": "The candidate matches the backend python developer profile perfectly."
                                    }
                                ]
                            }
                        }
                    ]
                },
            },
        )(),
    )

    detail_resp = client.get(f"/api/recommendations/{resume_id}/{j1_data['id']}")
    assert detail_resp.status_code == 200
    detail_data = detail_resp.json()
    assert detail_data["job_id"] == j1_data["id"]
    assert detail_data["score"] == 100
    assert detail_data["explanation"] == "The candidate matches the backend python developer profile perfectly."
